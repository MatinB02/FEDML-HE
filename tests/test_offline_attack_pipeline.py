import json
import sys
import tempfile
import types
import unittest
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import torch
from openpyxl import load_workbook

if "skimage.metrics" not in sys.modules:
    skimage_module = types.ModuleType("skimage")
    skimage_metrics_module = types.ModuleType("skimage.metrics")
    skimage_metrics_module.structural_similarity = lambda first, second, **_kwargs: float(
        max(-1.0, 1.0 - np.mean((np.asarray(first) - np.asarray(second)) ** 2))
    )
    sys.modules["skimage"] = skimage_module
    sys.modules["skimage.metrics"] = skimage_metrics_module

from Codes.attack_checkpoints import (
    AttackCheckpoint, AttackCheckpointWriter, CheckpointValidationError,
    construct_attacker_visible_state, state_schema,
)
from Codes.attack_selection import sample_class_matched_mia_data
from Codes.enums import DatasetType
from Codes.functions_mainAlg import build_mask_plan
from Codes.offline_attacks import (
    TaskFilter, execute_attack_run, plan_attack_run, reduce_attack_run,
)
from Models.DNNs.lenet5 import Lenet5


def _config(run_id="checkpoint-test"):
    return SimpleNamespace(
        moduleName="FEDML-HE",
        exp_id="offline-test",
        run_id=run_id,
        group="TEST",
        model=Lenet5,
        DB_dataset=DatasetType.MNIST,
        DB_samples_per_client=6,
        DB_nonIID_alpha=0.5,
        DB_nonIID=False,
        DB_forceCreate=False,
        num_clients=1,
        rounds=1,
        local_epochs=1,
        local_batch_size=2,
        eval_batch_size=4,
        mask_gradient_batch_size=2,
        encryption_ratio=0.1,
        aggregate_BN=False,
        temperature=4.0,
        seed=17,
        attack=True,
        attack_interval=1,
        attack_seed=101,
        mia_sample_size=4,
        mia_bootstrap_samples=5,
        ilrg_batch_size=2,
        ilrg_num_batches=2,
        ilrg_alpha=0.01,
        ilrg_mask_mode="partial",
        dlg_num_samples=2,
        dlg_num_restarts=2,
        dlg_iterations=2,
        dlg_learning_rate=0.01,
        dlg_optimizer="adam",
        dlg_objective="l2",
        dlg_tv_weight=0.0,
        dlg_early_stopping_patience=0,
        dlg_success_ssim=0.5,
        dlg_compute_lpips=False,
        dlg_known_label=True,
        ig_num_samples=2,
        ig_num_restarts=2,
        ig_iterations=2,
        ig_learning_rate=0.1,
        ig_tv_weight=0.0,
        ig_early_stopping_patience=0,
        ig_success_ssim=0.5,
        ig_compute_lpips=False,
        ig_known_label=True,
        inputShape=(12, 12, 1),
        classNum=2,
        resolved_args={"source": "unit-test"},
    )


def _data():
    rng = np.random.default_rng(11)
    train_x = rng.random((6, 12, 12, 1), dtype=np.float32)
    train_y = np.eye(2, dtype=np.float32)[[0, 1, 0, 1, 0, 1]]
    test_x = rng.random((6, 12, 12, 1), dtype=np.float32)
    test_y = np.eye(2, dtype=np.float32)[[1, 0, 1, 0, 1, 0]]
    return train_x, train_y, test_x, test_y


def _make_checkpoint(base: Path, encrypted_mode="partial") -> Path:
    torch.manual_seed(3)
    cfg = _config(run_id=f"checkpoint-{encrypted_mode}")
    model = Lenet5(num_classes=2, in_channels=1, input_height=12, input_width=12)
    initial = [value.detach().cpu().numpy().copy() for value in model.state_dict().values()]
    train_x, train_y, test_x, test_y = _data()
    writer = AttackCheckpointWriter(
        base_dir=base,
        project_root=Path(__file__).resolve().parents[1],
        cfg=cfg,
        model=model,
        initial_state=initial,
        client_data=[train_x],
        client_labels=[train_y],
        test_data=(test_x, test_y),
        dataset_train_data=(train_x, train_y),
    )
    actual = [value.copy() for value in initial]
    actual[0] = actual[0] + np.float32(0.25)
    if encrypted_mode == "all":
        encrypted = [np.ones_like(value, dtype=np.uint8) for value in initial]
        global_mask = None
    elif encrypted_mode == "none":
        encrypted = [np.zeros_like(value, dtype=np.uint8) for value in initial]
        global_mask = np.empty(0, dtype=np.int32)
    else:
        encrypted = [np.zeros_like(value, dtype=np.uint8) for value in initial]
        encrypted[0].reshape(-1)[::3] = 1
        flat_mask = np.concatenate([value.reshape(-1) for value in encrypted])
        global_mask = np.flatnonzero(flat_mask).astype(np.int32)
    plaintext = [1 - value for value in encrypted]
    exposed_before = np.concatenate([value.reshape(-1) for value in initial]).astype(np.float32)
    exposed_after = np.full_like(exposed_before, 9.0)
    writer.prepare_round(
        round_index=0,
        global_entering=initial,
        exposed_flat_before_update=exposed_before,
        encrypted_masks=encrypted,
        plaintext_masks=plaintext,
        global_mask=global_mask,
        mask_plan=build_mask_plan(encrypted, plaintext),
        sensitivity_maps=[[
            np.zeros_like(parameter.detach().cpu().numpy(), dtype=np.float32)
            for parameter in model.parameters()
        ]],
        client_states=[actual],
        client_weights=[1.0],
        client_dataset_sizes=[len(train_x)],
        transmission_metadata=[{"plaintext_payload_bytes": 1, "encrypted_payload_bytes": 2}],
    )
    writer.finalize_round(
        global_after_aggregation=actual,
        exposed_flat_after_update=exposed_after,
        post_aggregation_metadata={"timing": "after aggregation"},
    )
    writer.finalize_experiment()
    return writer.run_dir


class OfflineCheckpointTests(unittest.TestCase):
    def test_round_trip_namespaces_hashes_masks_and_correct_timing(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = _make_checkpoint(Path(directory), "partial")
            checkpoint = AttackCheckpoint(run_dir)
            artifacts = checkpoint.load_artifacts(0, 0)

            actual = artifacts.evaluation_oracle["actual_post_local_state"]
            attacker = artifacts.server_visible["attacker_visible_state"]
            encrypted = artifacts.server_visible["encrypted_masks"]
            plaintext = artifacts.server_visible["plaintext_masks"]
            entering = artifacts.server_visible["global_entering_state"]
            sensitivity = artifacts.evaluation_oracle["sensitivity_map"]
            expected_model = Lenet5(
                num_classes=2, in_channels=1, input_height=12, input_width=12
            )
            self.assertEqual(
                len(sensitivity), len(list(expected_model.parameters()))
            )
            for value, parameter in zip(sensitivity, expected_model.parameters()):
                np.testing.assert_array_equal(
                    value, np.zeros_like(parameter.detach().cpu().numpy())
                )
            sensitivity_metadata = artifacts.evaluation_oracle[
                "sensitivity_map_metadata"
            ]
            self.assertEqual(sensitivity_metadata["namespace"], "evaluation_oracle")
            self.assertFalse(sensitivity_metadata["available_to_server"])
            self.assertEqual(sensitivity_metadata["aggregation_weights"], [1.0])
            for actual_value, attacker_value, entering_value, enc, plain in zip(
                    actual, attacker, entering, encrypted, plaintext):
                np.testing.assert_array_equal(attacker_value[plain.astype(bool)], actual_value[plain.astype(bool)])
                np.testing.assert_array_equal(attacker_value[enc.astype(bool)], entering_value[enc.astype(bool)])
                self.assertEqual(actual_value.dtype, attacker_value.dtype)
                self.assertEqual(enc.dtype, np.uint8)
                self.assertEqual(actual_value.shape, attacker_value.shape)

            selections = artifacts.evaluation_oracle["selections"]
            train_x, train_y, test_x, test_y = _data()
            expected_mia = sample_class_matched_mia_data(
                (train_x, train_y), (test_x, test_y), 4, 101
            )
            for actual_selection, expected_selection in zip(
                    (
                        selections["mia_member_inputs"], selections["mia_member_labels"],
                        selections["mia_nonmember_inputs"], selections["mia_nonmember_labels"],
                    ),
                    expected_mia):
                np.testing.assert_array_equal(actual_selection, expected_selection)
            self.assertEqual(len(selections["mia_member_inputs"]), 4)
            np.testing.assert_array_equal(
                np.bincount(selections["mia_member_classes"], minlength=2),
                np.bincount(selections["mia_nonmember_classes"], minlength=2),
            )
            self.assertEqual(len(selections["ilrg_batch_ids"]), 2)
            self.assertEqual(selections["dlg_restart_seeds"].shape, (2, 2))
            self.assertEqual(selections["ig_restart_seeds"].shape, (2, 2))
            expected_batch_ids = np.sort(
                np.random.default_rng(101).choice(3, size=2, replace=False)
            )
            np.testing.assert_array_equal(selections["ilrg_batch_ids"], expected_batch_ids)
            expected_samples = np.random.default_rng(101).choice(6, size=2, replace=False)
            np.testing.assert_array_equal(selections["dlg_sample_ids"], expected_samples)
            np.testing.assert_array_equal(selections["ig_sample_ids"], expected_samples)

    def test_state_schema_preserves_batchnorm_buffers_and_parameter_order(self):
        model = torch.nn.Sequential(
            torch.nn.Conv2d(1, 2, 1), torch.nn.BatchNorm2d(2), torch.nn.Flatten(),
            torch.nn.Linear(8, 2),
        )
        schema, trainable = state_schema(model)
        self.assertEqual([entry["name"] for entry in schema], list(model.state_dict()))
        classifications = {entry["name"]: entry["classification"] for entry in schema}
        self.assertEqual(classifications["1.running_mean"], "buffer")
        self.assertEqual(classifications["1.running_var"], "buffer")
        self.assertEqual(classifications["1.num_batches_tracked"], "buffer")
        self.assertNotIn("1.running_mean", [entry["name"] for entry in trainable])

    def test_zero_and_full_visibility_use_compact_global_mask_representations(self):
        with tempfile.TemporaryDirectory() as directory:
            none_dir = _make_checkpoint(Path(directory), "none")
            all_dir = _make_checkpoint(Path(directory), "all")
            none_checkpoint = AttackCheckpoint(none_dir)
            all_checkpoint = AttackCheckpoint(all_dir)
            self.assertEqual(
                none_checkpoint.load_round_manifest(0)["consensus_mask"]["representation"],
                "none",
            )
            self.assertEqual(
                all_checkpoint.load_round_manifest(0)["consensus_mask"]["representation"],
                "all",
            )
            none_artifacts = none_checkpoint.load_artifacts(0, 0)
            all_artifacts = all_checkpoint.load_artifacts(0, 0)
            for actual, attacker in zip(
                    none_artifacts.evaluation_oracle["actual_post_local_state"],
                    none_artifacts.server_visible["attacker_visible_state"]):
                np.testing.assert_array_equal(actual, attacker)
            for entering, attacker in zip(
                    all_artifacts.server_visible["global_entering_state"],
                    all_artifacts.server_visible["attacker_visible_state"]):
                np.testing.assert_array_equal(entering, attacker)

    def test_incomplete_checkpoint_and_round_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = _make_checkpoint(Path(directory), "none")
            (run_dir / "CHECKPOINT_COMPLETE").unlink()
            with self.assertRaisesRegex(CheckpointValidationError, "incomplete"):
                AttackCheckpoint(run_dir)
        with tempfile.TemporaryDirectory() as directory:
            run_dir = _make_checkpoint(Path(directory), "none")
            (run_dir / "round_0000" / "ROUND_COMPLETE").unlink()
            checkpoint = AttackCheckpoint(run_dir)
            with self.assertRaisesRegex(CheckpointValidationError, "incomplete"):
                checkpoint.load_round_manifest(0)

    def test_task_planning_execution_resume_reduction_and_config_identity(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            checkpoint_dir = _make_checkpoint(root / "checkpoints", "none")
            output_root = root / "offline"
            overrides = {
                "attacks": ["dlg"],
                "dlg_num_samples": 1,
                "dlg_num_restarts": 2,
                "dlg_iterations": 1,
                "dlg_known_label": True,
            }
            run_root, manifest = plan_attack_run(
                checkpoint_dir, overrides=overrides, filters=TaskFilter(),
                output_root=output_root,
            )
            self.assertEqual(len(manifest["tasks"]), 2)
            outcome = execute_attack_run(
                run_root, manifest, devices=["cpu"], workers_per_device=2,
            )
            self.assertEqual(outcome["succeeded"], 2)
            resumed = execute_attack_run(
                run_root, manifest, devices=["cpu"], workers_per_device=1,
                resume=True,
            )
            self.assertEqual(resumed["skipped"], 2)

            corrupt_task = manifest["tasks"][0]
            corrupt_result = (
                run_root / "tasks" / corrupt_task["attack"] / corrupt_task["task_id"] /
                "result.json"
            )
            corrupt_result.write_text("{\"status\": \"succeeded\"}", encoding="utf-8")
            repaired = execute_attack_run(
                run_root, manifest, devices=["cpu"], workers_per_device=1,
                resume=True,
            )
            self.assertEqual(repaired["skipped"], 1)
            self.assertEqual(repaired["succeeded"], 1)
            report = reduce_attack_run(run_root, manifest)
            self.assertEqual(report["task_counts"]["succeeded"], 2)
            self.assertTrue((run_root / "records.jsonl").is_file())
            self.assertTrue((run_root / "report.json").is_file())
            workbook = load_workbook(run_root / "offline_attacks.xlsx", read_only=True)
            self.assertIn("Attack_DLG", workbook.sheetnames)
            self.assertIn("Attack_DLG_Summary", workbook.sheetnames)
            workbook.close()
            first_report = json.loads((run_root / "report.json").read_text(encoding="utf-8"))
            reduce_attack_run(run_root, manifest)
            second_report = json.loads((run_root / "report.json").read_text(encoding="utf-8"))
            self.assertEqual(
                first_report["client_summaries"], second_report["client_summaries"]
            )

            changed_root, changed_manifest = plan_attack_run(
                checkpoint_dir,
                overrides={**overrides, "dlg_learning_rate": 0.02},
                filters=TaskFilter(), output_root=output_root,
            )
            self.assertNotEqual(manifest["attack_run_id"], changed_manifest["attack_run_id"])
            self.assertNotEqual(run_root, changed_root)


if __name__ == "__main__":
    unittest.main()
