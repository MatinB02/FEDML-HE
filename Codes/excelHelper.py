# -*- coding: utf-8 -*-

import json
import os
import tempfile
from pathlib import Path
from types import SimpleNamespace

import numpy as np
from openpyxl import Workbook
from openpyxl.utils.exceptions import InvalidFileException
from Codes.enums import Config

time_names = {"time_train": 1,
              "time_getLogits": 2,
              "time_getWeights": 3,
              "time_weightCalc": 5,
              "time_aggregation": 6,
              "time_masking": 7,
              "time_encryption": 8,
              "time_reconstructing": 9,
              "time_ClientMaskProposal": 10,
              "time_sensitivityMapsAggregation": 11,
              "time_MaskDecryption": 12,
              "time_MaskGen": 13,
              "time_checkpoint": 14}

_pending_records = {}


def _jsonl_path(excel_addr):
    return Path(str(excel_addr) + ".jsonl")


def _json_compatible(value):
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, dict):
        return {str(key): _json_compatible(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_compatible(item) for item in value]
    return value


def _new_metric_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Accuracy MAX"
    for sheet_name in (
            "Train Accuracy MAX", "Test Accuracy", "Train Accuracy", "Loss",
            "global_test_acc", "sensitivityScore", "totalSize_encrypted",
            "totalSize_plaintext", "totalSize_original", "totalSize_Sum",
            "size_plaintext", "size_cyphertext", "size_sensitivity_map"):
        wb.create_sheet(sheet_name)
    global_ws = wb["global_test_acc"]
    global_ws.cell(row=1, column=1).value = "Round"
    global_ws.cell(row=1, column=2).value = "Global Test Accuracy"
    global_ws.freeze_panes = "A2"
    ws = wb.create_sheet("Times")
    for key, column in time_names.items():
        ws.cell(row=1, column=column).value = str(key).split("_")[1]
    ws.freeze_panes = "A2"
    return wb
def _column_map(keys):
    return {key: column for column, key in enumerate(keys, start=1)}


_mia_metric_suffixes = (
    "AUC",
    "AUC_CI_Lower",
    "AUC_CI_Upper",
    "BootstrapSamples",
    "BootstrapScheme",
    "OracleBestAcc",
    "BalancedAcc",
    "OptimalThreshold",
    "Precision",
    "Recall",
    "F1",
    "Specificity",
    "PrivacyAdvantage",
    "TPRAtFPR1pct",
    "TPRAtFPR5pct",
    "TPRAtFPR10pct",
    "AttackTime",
    "SharedFeatureExtractionTime",
    "SignalEvaluationTime",
    "NumMember",
    "NumNonmember",
    "ThresholdNote",
    "FPR",
    "TPR",
)

_mia_keys = ["round", "client", "MIA_TotalTime"]
for _signal in ("loss", "entropy", "modified_entropy"):
    for _suffix in _mia_metric_suffixes:
        _mia_keys.append(f"{_signal}_{_suffix}")

mia_attack_metrics = _column_map(_mia_keys)

ilrg_attack_metrics = _column_map((
    "round", "client", "batch_id", "batch_start", "batch_end", "batch_size",
    "attack_available", "unavailable_reason", "mask_mode", "attack_scope",
    "alpha", "true_labels", "predicted_labels", "true_counts",
    "continuous_counts", "predicted_counts", "label_existence_accuracy",
    "label_number_accuracy", "instance_recall", "count_mae",
    "normalized_count_l1", "count_cosine_similarity", "exact_count_vector",
    "label_precision", "label_recall", "label_f1", "attack_time",
    "visible_gradient_fraction", "final_weight_visible_fraction",
    "final_bias_visible_fraction", "usable_bias_equations",
    "recovered_embedding_classes", "mean_embedding_visible_fraction",
    "system_rank", "system_condition_number", "residual_l2",
    "target_gradient_state", "inference_model_state", "final_layer_name",
))

ilrg_summary_metrics = _column_map((
    "round", "client", "num_batches", "num_available", "availability_rate",
    "total_attack_time",
    "label_existence_accuracy_mean", "label_existence_accuracy_std",
    "label_number_accuracy_mean", "label_number_accuracy_std",
    "instance_recall_mean", "instance_recall_std",
    "count_mae_mean", "count_mae_std",
    "normalized_count_l1_mean", "normalized_count_l1_std",
    "count_cosine_similarity_mean", "count_cosine_similarity_std",
    "exact_count_vector_mean", "exact_count_vector_std",
    "label_precision_mean", "label_precision_std",
    "label_recall_mean", "label_recall_std", "label_f1_mean", "label_f1_std",
    "attack_time_mean", "attack_time_std",
    "visible_gradient_fraction_mean", "visible_gradient_fraction_std",
    "final_weight_visible_fraction_mean", "final_weight_visible_fraction_std",
    "final_bias_visible_fraction_mean", "final_bias_visible_fraction_std",
    "usable_bias_equations_mean", "usable_bias_equations_std",
    "recovered_embedding_classes_mean", "recovered_embedding_classes_std",
    "mean_embedding_visible_fraction_mean", "mean_embedding_visible_fraction_std",
))

dlg_attack_metrics = _column_map((
    "round", "client", "sample_id", "true_label", "inferred_label",
    "label_inference_available", "label_inference_method", "known_label",
    "idlg_inferred_label", "idlg_label_inference_success", "label_accuracy",
    "success_rate", "mse",
    "psnr", "ssim", "cosine_sim", "lpips", "attack_time", "best_loss",
    "best_gradient_loss", "best_iteration", "best_restart",
    "best_restart_seed", "final_loss", "iterations_run", "num_restarts",
    "visible_gradient_fraction", "optimizer", "objective", "learning_rate",
    "tv_weight", "success_ssim_threshold", "attack_scope",
))

dlg_summary_metrics = _column_map((
    "round", "client", "num_samples", "total_attack_time",
    "success_rate_mean", "success_rate_std",
    "label_accuracy_mean", "label_accuracy_std",
    "idlg_label_inference_success_mean", "idlg_label_inference_success_std",
    "mse_mean", "mse_std", "psnr_mean", "psnr_std",
    "ssim_mean", "ssim_std", "cosine_sim_mean", "cosine_sim_std",
    "lpips_mean", "lpips_std", "attack_time_mean", "attack_time_std",
    "visible_gradient_fraction_mean", "visible_gradient_fraction_std",
))

ig_attack_metrics = _column_map((
    "round", "client", "sample_id", "true_label", "inferred_label",
    "label_inference_available", "label_inference_method", "known_label",
    "idlg_inferred_label", "idlg_label_inference_success", "label_accuracy",
    "success_rate", "mse", "psnr", "ssim", "cosine_sim", "lpips",
    "attack_time", "best_loss", "best_gradient_loss", "best_iteration",
    "best_restart", "best_restart_seed", "final_loss", "iterations_run",
    "num_restarts", "visible_gradient_fraction", "optimizer", "objective",
    "initialization", "learning_rate", "lr_decay_gamma",
    "lr_decay_milestones", "tv_weight", "success_ssim_threshold",
    "attack_scope",
))

ig_summary_metrics = _column_map((
    "round", "client", "num_samples", "total_attack_time",
    "success_rate_mean", "success_rate_std",
    "label_accuracy_mean", "label_accuracy_std",
    "idlg_label_inference_success_mean", "idlg_label_inference_success_std",
    "mse_mean", "mse_std", "psnr_mean", "psnr_std",
    "ssim_mean", "ssim_std", "cosine_sim_mean", "cosine_sim_std",
    "lpips_mean", "lpips_std", "attack_time_mean", "attack_time_std",
    "visible_gradient_fraction_mean", "visible_gradient_fraction_std",
))

generic_attack_metrics = _column_map((
    "round", "client", "success_rate", "cosine_sim", "mse", "psnr",
    "ssim", "lpips", "attack_time",
))


def _attack_schema(sheet_name):
    if sheet_name == "Attack_MIA":
        return mia_attack_metrics
    if sheet_name == "Attack_iLRG":
        return ilrg_attack_metrics
    if sheet_name == "Attack_iLRG_Summary":
        return ilrg_summary_metrics
    if sheet_name == "Attack_DLG":
        return dlg_attack_metrics
    if sheet_name == "Attack_DLG_Summary":
        return dlg_summary_metrics
    if sheet_name == "Attack_IG":
        return ig_attack_metrics
    if sheet_name == "Attack_IG_Summary":
        return ig_summary_metrics
    return generic_attack_metrics


def _write_headers(ws, column_map):
    for key, column in column_map.items():
        ws.cell(row=1, column=column).value = key
    ws.freeze_panes = "A2"


def write_offline_attack_workbook(records_by_sheet, output_path, provenance=None):
    """Write all reduced offline attack records in one atomic workbook save."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    provenance_sheet = workbook.active
    provenance_sheet.title = "Provenance"
    provenance_sheet.append(["field", "value"])
    provenance_sheet.freeze_panes = "A2"
    for key, value in sorted((provenance or {}).items()):
        if isinstance(value, (dict, list, tuple)):
            value = json.dumps(_json_compatible(value), sort_keys=True, separators=(",", ":"))
        provenance_sheet.append([key, value])

    for sheet_name, records in records_by_sheet.items():
        worksheet = workbook.create_sheet(title=sheet_name)
        schema = _attack_schema(sheet_name)
        _write_headers(worksheet, schema)
        for row_index, record in enumerate(records, start=2):
            for key, value in record.items():
                if key not in schema:
                    continue
                if isinstance(value, (dict, list, tuple)):
                    value = json.dumps(
                        _json_compatible(value), sort_keys=True, separators=(",", ":")
                    )
                worksheet.cell(row=row_index, column=schema[key]).value = value

    fd, temp_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}.", suffix=".xlsx", dir=output_path.parent
    )
    os.close(fd)
    try:
        workbook.save(temp_name)
        os.replace(temp_name, output_path)
    except BaseException:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass
        raise


def create(cfg: Config):
    if cfg.excelAddr is None:
        return
    metrics_path = _jsonl_path(cfg.excelAddr)
    metrics_path.parent.mkdir(parents=True, exist_ok=True)
    metrics_path.write_text("", encoding="utf-8")
    _pending_records[metrics_path] = []


def update(cfg: Config, dataDic):
    if cfg.excelAddr is None:
        raise InvalidFileException("cfg.excelAddr is none")
    metrics_path = _jsonl_path(cfg.excelAddr)
    record = {
        "context": {
            "currentRound": getattr(cfg, "currentRound", 0),
            "currentEdge": getattr(cfg, "currentEdge", 0),
            "currentEpoch": getattr(cfg, "currentEpoch", 0),
            "num_clients": getattr(cfg, "num_clients", 1),
        },
        "metrics": _json_compatible(dataDic),
    }
    _pending_records.setdefault(metrics_path, []).append(record)


def _flush_path(metrics_path):
    records = _pending_records.get(metrics_path, [])
    if not records:
        return
    metrics_path.parent.mkdir(parents=True, exist_ok=True)
    with metrics_path.open("a", encoding="utf-8") as metrics_file:
        for record in records:
            metrics_file.write(json.dumps(record, separators=(",", ":")) + "\n")
    records.clear()


def flush(cfg: Config):
    """Append the current round's buffered metrics to its JSONL journal."""
    if cfg.excelAddr is None:
        raise InvalidFileException("cfg.excelAddr is none")
    _flush_path(_jsonl_path(cfg.excelAddr))


def _apply_update(wb, context, dataDic):
    cfg = SimpleNamespace(**context)

    for sheetName, value in dataDic.items():
        sheet_str = str(sheetName)

        # ---------- TIME SHEETS ----------
        if sheet_str.__contains__("time"):
            target_sheet_name = "Times"

            # Create sheet if it doesn't exist
            if target_sheet_name in wb.sheetnames:
                ws = wb[target_sheet_name]
            else:
                ws = wb.create_sheet(title=target_sheet_name)

            # Get column for this time metric
            col_idx = time_names[sheetName]

            # Find next empty row in that column
            currentRow = __fullRowCount(ws, colNum=col_idx)
            ws.cell(row=currentRow + 1, column=col_idx).value = value

        # ---------- CLIENT/SERVER PAYLOAD SIZE SHEETS ----------
        elif sheet_str in ("size_plaintext", "size_cyphertext", "size_sensitivity_map"):
            if sheet_str in wb.sheetnames:
                ws = wb[sheet_str]
            else:
                ws = wb.create_sheet(title=sheet_str)
            ws.cell(row=1, column=cfg.currentEdge + 1).value = int(value)

        elif sheet_str.startswith("size_server_"):
            target_sheet_name = sheet_str.replace("size_server_", "size_", 1)
            if target_sheet_name in wb.sheetnames:
                ws = wb[target_sheet_name]
            else:
                ws = wb.create_sheet(title=target_sheet_name)
            ws.cell(row=1, column=cfg.num_clients + 1).value = int(value)

        # ---------- LEGACY TOTAL SIZE SHEETS ----------
        elif sheet_str.__contains__("totalSize"):
            target_sheet_name = sheet_str  # uses the key itself as sheet name
            if target_sheet_name in wb.sheetnames:
                ws = wb[target_sheet_name]
            else:
                ws = wb.create_sheet(title=target_sheet_name)

            ws.cell(row=1, column=cfg.currentEdge + 1).value = value

        # ---------- SENSITIVITY SCORE SHEET ----------
        elif sheet_str.__contains__("sensitivityScore"):
            target_sheet_name = "sensitivityScore"
            if target_sheet_name in wb.sheetnames:
                ws = wb[target_sheet_name]
            else:
                ws = wb.create_sheet(title=target_sheet_name)
            ws.cell(row=1, column=cfg.currentEdge + 1).value = value

        # ---------- GLOBAL ACCURACY (one record per communication round) ----------
        elif sheet_str == "global_test_acc":
            if sheet_str in wb.sheetnames:
                ws = wb[sheet_str]
            else:
                ws = wb.create_sheet(title=sheet_str)
            ws.cell(row=1, column=1).value = "Round"
            ws.cell(row=1, column=2).value = "Global Test Accuracy"
            ws.freeze_panes = "A2"
            current_row = max(2, ws.max_row + 1)
            ws.cell(row=current_row, column=1).value = int(cfg.currentRound)
            ws.cell(row=current_row, column=2).value = float(value)

        # ---------- ATTACK SHEETS ----------
        elif sheet_str.__contains__("Attack"):
            target_sheet_name = sheet_str
            if target_sheet_name in wb.sheetnames:
                ws = wb[target_sheet_name]
            else:
                ws = wb.create_sheet(title=target_sheet_name)

            attack_metrics = _attack_schema(target_sheet_name)

            # We determine the row to write based on the first column tracking metrics
            currentRow = __fullRowCount(ws, colNum=1) + 1
            if currentRow == 1:
                currentRow = 2 # Row 1 is reserved exclusively for headers

            # Write headers on row 1 explicitly mapping out our attack keys
            _write_headers(ws, attack_metrics)

            # Write values to the specified columns
            if isinstance(value, dict):
                for col_title, col_value in value.items():
                    if col_title in attack_metrics:
                        col_idx = attack_metrics[col_title]
                        ws.cell(row=currentRow, column=col_idx).value = col_value
            else:
                # If a fallback single scalar value gets hit
                ws.cell(row=currentRow, column=1).value = value

        # ---------- DEFAULT / EPOCH-BASED SHEETS ----------
        else:
            target_sheet_name = sheet_str
            if target_sheet_name in wb.sheetnames:
                ws = wb[target_sheet_name]
            else:
                ws = wb.create_sheet(title=target_sheet_name)

            ws.cell(row=cfg.currentEpoch + 1, column=cfg.currentEdge + 1).value = value

def _workbook_from_jsonl(metrics_path):
    _flush_path(metrics_path)
    wb = _new_metric_workbook()
    if not metrics_path.exists():
        return wb
    with metrics_path.open("r", encoding="utf-8") as metrics_file:
        for line_number, line in enumerate(metrics_file, start=1):
            if not line.strip():
                continue
            try:
                record = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(
                    f"Invalid metrics JSONL at {metrics_path}:{line_number}"
                ) from exc
            _apply_update(wb, record["context"], record["metrics"])
    return wb


def combineExcels(cfg: Config, baseAddr, targetSaveAddr, sheetList, newDataDic=None, newDataType=None):
    # combine excels
    TYPE_MAP = {
        "str": str,
        "int": int,
        "float": float,
        "np.int32": np.int32,
        "int32": np.int32,
        "np.float32": np.float32,
        "float32": np.float32,
    }
    wb_result = Workbook()
    ws_result = wb_result.active
    ws_result.title = sheetList[0]
    for sheet_index, sheet_name in enumerate(sheetList):
        if sheet_index > 0:
            ws_result = wb_result.create_sheet(sheet_name)
        else:
            ws_result = wb_result[sheet_name]

        if "Time" in str(sheet_name):
            for key, column in time_names.items():
                ws_result.cell(row=1, column=column).value = str(key).split("_")[1]
        elif "Attack" in str(sheet_name):
            _write_headers(ws_result, _attack_schema(str(sheet_name)))
        elif str(sheet_name) == "global_test_acc":
            ws_result.cell(row=1, column=1).value = "Round"
            ws_result.cell(row=1, column=2).value = "Global Test Accuracy"
            ws_result.freeze_panes = "A2"
        else:
            for client_index in range(cfg.num_clients):
                ws_result.cell(1, client_index + 1, "Client %d" % (client_index + 1))
            if str(sheet_name) in (
                    "size_plaintext", "size_cyphertext", "size_sensitivity_map"):
                ws_result.cell(row=1, column=cfg.num_clients + 1, value="Server")
    resultRowCount = [2] * len(sheetList)
    iterationList = list(range(cfg.rounds))  # [0, 1, 2, 3, ]
    for iter in iterationList:
        metrics_path = _jsonl_path(baseAddr % iter)
        print(metrics_path)
        wb = _workbook_from_jsonl(metrics_path)
        for shIndex in range(len(sheetList)):
            sheetName = sheetList[shIndex]
            if sheetName in wb.sheetnames:
                if str(sheetName).__contains__("Time"):
                    ws = wb[sheetName]
                    ws_result = wb_result[sheetName]
                    if resultRowCount[shIndex] == 2:
                        for keys in time_names.keys():
                            ws_result.cell(row=1, column=time_names[keys]).value = str(keys).split("_")[1]
                    times_mean = {a: [] for a in time_names}
                    for j in range(cfg.num_clients):
                        for keys in time_names.keys():
                            if ws.cell(row=j + 2, column=time_names[keys]).value is not None:
                                times_mean[keys].append(np.float32(ws.cell(row=j + 2, column=time_names[keys]).value))
                    for keys in time_names.keys():
                        values = times_mean[keys]
                        ws_result.cell(
                            row=resultRowCount[shIndex], column=time_names[keys]
                        ).value = float(np.mean(values)) if values else None

                    resultRowCount[shIndex] += 1

                elif str(sheetName).__contains__("Attack"):
                    ws = wb[sheetName]
                    ws_result = wb_result[sheetName]
                    attack_metrics = _attack_schema(str(sheetName))
                    if resultRowCount[shIndex] == 2:
                        _write_headers(ws_result, attack_metrics)

                    # A MIA row may have no value in the first generic attack
                    # column, so detect data across every registered metric.
                    for row_index in range(2, ws.max_row + 1):
                        if not any(
                                ws.cell(row=row_index, column=column).value is not None
                                for column in attack_metrics.values()):
                            continue
                        for keys in attack_metrics.keys():
                            ws_result.cell(row=resultRowCount[shIndex], column=attack_metrics[keys]).value = \
                                ws.cell(row=row_index, column=attack_metrics[keys]).value
                        resultRowCount[shIndex] += 1

                elif str(sheetName) == "global_test_acc":
                    ws = wb[sheetName]
                    ws_result = wb_result[sheetName]
                    for row_index in range(2, ws.max_row + 1):
                        if ws.cell(row=row_index, column=2).value is None:
                            continue
                        ws_result.cell(
                            row=resultRowCount[shIndex], column=1
                        ).value = ws.cell(row=row_index, column=1).value
                        ws_result.cell(
                            row=resultRowCount[shIndex], column=2
                        ).value = ws.cell(row=row_index, column=2).value
                        resultRowCount[shIndex] += 1

                elif (str(sheetName) in (
                        "size_plaintext", "size_cyphertext", "size_sensitivity_map") or
                      str(sheetName).__contains__("sensitivityScore") or
                      str(sheetName).__contains__("totalSize")):
                    ws = wb[sheetName]
                    ws_result = wb_result[sheetName]
                    column_count = (
                        cfg.num_clients + 1
                        if str(sheetName) in (
                            "size_plaintext", "size_cyphertext", "size_sensitivity_map"
                        )
                        else cfg.num_clients
                    )
                    for j in range(column_count):
                        ws_result.cell(row=resultRowCount[shIndex], column=j + 1).value = ws.cell(row=1, column=j + 1).value
                    resultRowCount[shIndex] += 1
                else:
                    ws = wb[sheetName]
                    ws_result = wb_result[sheetName]
                    for source_row in range(1, ws.max_row + 1):
                        if not any(
                                ws.cell(row=source_row, column=column).value is not None
                                for column in range(1, cfg.num_clients + 1)):
                            continue
                        for j in range(cfg.num_clients):
                            ws_result.cell(
                                row=resultRowCount[shIndex], column=j + 1
                            ).value = ws.cell(
                                row=source_row, column=j + 1
                            ).value
                        resultRowCount[shIndex] += 1
        wb.close()

    if newDataDic is not None:
        for pageTitle in newDataDic.keys():
            data = newDataDic[pageTitle]
            for c in range(cfg.num_clients):
                ws_result = wb_result.create_sheet(pageTitle + "_C%d" % c)
                roundsCount = len(data)
                for row in range(roundsCount):
                    clientData = data[row][c]
                    for col in range(len(clientData)):
                        convertType = TYPE_MAP[newDataType[pageTitle]]
                        ws_result.cell(row + 1, col + 1, convertType(clientData[col]))

    target_path = Path(targetSaveAddr + '.xlsx')
    target_path.parent.mkdir(parents=True, exist_ok=True)
    wb_result.save(target_path)
    wb_result.close()


def __fullRowCount(sheet, colNum):
    rowIndex = 1
    while sheet.cell(row=rowIndex, column=colNum).value is not None:
        rowIndex += 1
    return rowIndex - 1
